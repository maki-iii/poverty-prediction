from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).parent
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.72)
sec.left_margin = sec.right_margin = Inches(.82)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for name, size in [('Title',30),('Subtitle',14),('Heading 1',21),('Heading 2',14),('Heading 3',12)]:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0,0,0)
    st.paragraph_format.space_after = Pt(9)
    st.paragraph_format.space_before = Pt(12)
    st.paragraph_format.keep_with_next = True
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run('PPS User Manual  |  ')
r.font.size = Pt(9)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
doc.core_properties.title = 'Poverty Prediction System User Manual'
doc.core_properties.subject = 'NCR FIES workflow and system operation'
doc.core_properties.author = 'Poverty Prediction System'

def p(s,style=None): return doc.add_paragraph(s,style)
def h(s): doc.add_heading(s,2)
def steps(items):
    for i,s in enumerate(items,1):
        q=p(f'{i}.  {s}'); q.paragraph_format.left_indent=Inches(.16); q.paragraph_format.first_line_indent=Inches(-.16)
def bullets(items):
    for s in items: p(s,'List Bullet')
def note(s):
    q=p('');q.add_run('Note. ').bold=True;q.add_run(s)
def page(n,title):
    doc.add_page_break()
    q=doc.add_heading(f'{n}  {title}',1)
    b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(n));b.set(qn('w:name'),f'section{n}');q._p.insert(0,b)
    e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(n));q._p.append(e)
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1, cols=len(headers));t.autofit=False
    for i,s in enumerate(headers):t.rows[0].cells[i].text=s
    for row in rows:
        c=t.add_row().cells
        for i,s in enumerate(row):c[i].text=str(s)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[ci])
            cp=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'002366' if ri==0 else ('F2F5F8' if ri%2 else 'FFFFFF'));cp.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');margins.append(el)
            cp.append(margins)
            for q in c.paragraphs:
                q.paragraph_format.space_after=Pt(2);q.paragraph_format.space_before=Pt(2)
                for r in q.runs:
                    r.font.size=Pt(10)
                    if ri==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    t._tbl.tblPr.append(borders)
    p('')
    return t
def code(s):
    q=p(s);q.paragraph_format.space_after=Pt(9)
    for r in q.runs:r.font.name='Consolas';r.font.size=Pt(9)

p('Poverty Prediction System','Title')
p('User Manual','Subtitle')
p('NCR FIES data preparation, forecasting and poverty prediction')
p('Version 1.0  |  22 October 2026')
h('About this manual')
p('This manual explains the processes available in the current administrator interface. It is for administrators preparing datasets, super administrators managing accounts, evaluators reviewing results, and the operator running the local Docker services.')
p('The active setup uses FIES data and interpolation, with three model inputs: average family income, average family expenditure, and average family size. APIS and inflation are not active inputs in this setup.')
h('Where to open the system')
p('Public frontend: https://poverty-prediction-system-five.vercel.app')
p('On the operator computer: http://localhost:8080')
p('The public frontend depends on the operator computer, Docker services and ngrok tunnel being online. A page opening successfully does not by itself confirm that the backend is available.')
h('The normal working order')
p('Sign in → Upload → Process → Merge → Review Data → Forecast → Train Models → Review Evaluation → Analyze and visualize results.')
note('Use the procedures in order when adding a new survey year. Viewing results does not require uploading or retraining. Do not delete existing records simply to demonstrate the system.')
h('Document scope')
p('Instructions are based on the current project source and configuration. Button names are written as they appear in the application. Numerical examples are illustrative, not official PSA statistics. Keep original files and a separate record of each official threshold and incidence source.')

page(1,'Contents and quick start')
sections=[('Access and navigation',2),('Prepare the source data',3),('Upload one or many datasets',4),('Process a dataset',5),('Merge processed datasets',6),('Review data and interpolation',7),('Run the three year forecast',8),('Train models and predict poverty',9),('Read the model evaluation',10),('Analysis and visualization',11),('Manage users and audit logs',12),('Cancel work and resolve errors',13),('Start and stop evaluator access',14),('Maintain deployment and preserve data',15)]
for label,n in sections:
    q=p('')
    link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),f'section{n}')
    r=OxmlElement('w:r');tx=OxmlElement('w:t');tx.text=f'{n}   {label}';r.append(tx);link.append(r);q._p.append(link)
h('First successful run')
steps(['Ask the operator to start the services, then sign in with your assigned account.','Upload a valid FIES file and process it using the correct survey year, annual NCR poverty threshold per person, and official family poverty incidence.','Review the processed values and merge the approved dataset. Verify its year and source in Data.','Run Forecast, then use Model Training to train and generate predictions. Wait for completion before switching modules.','Check the forecast and prediction years, compare model metrics, and review Analysis and Visualization.'])
note('If data and results already exist, begin with Data and the results pages. Repeat processing only when adding or correcting a dataset.')

page(2,'Access and navigation')
h('Sign in')
steps(['Open the public frontend address, or localhost:8080 on the operator computer.','On PPS Admin Portal, enter your Username and Password. Use the eye icon if you need to check the password you typed.','Click Secure Sign In. Wait for the Dashboard to open.','If access is denied, check the credentials and contact the super administrator. The active admin login screen does not provide a password recovery button.'])
h('Find a module')
table(['Module','Use it to'],[
('Dashboard','Review the poverty timeline, summary cards, recent activity and model performance.'),
('Dataset Management','Upload files, process them and merge approved values.'),
('Data','Inspect actual and interpolated training history.'),
('Model','Open Model Training or Model Evaluation.'),
('Forecasting','Generate and inspect the next three years of indicator inputs.'),
('Analysis and Visualization','Interpret feature importance, trends and indicator relationships.'),
('System','Super administrators access User Management and Audit Logs.')],[1.6,5.0])
h('Use the Dashboard')
p('Check the available history and latest predictions before starting work. Hover over chart points to read their year and value. The Activity area shows recent actions; it is not the full audit log. Model Performance summarizes saved evaluation results.')
h('Sign out')
p('Open the account menu at the bottom of the sidebar and choose Logout. Confirm with Log Out. If an operation is active, the leave-page warning appears first. Complete the operation or follow the cancellation procedure in section 13.')

page(3,'Prepare the source data')
p('Prepare a household-level FIES file in CSV, XLSX or XLS format. Keep an untouched copy. Use consistent definitions and units across survey years; do not combine records from different years into one file for a single processing run.')
table(['Required information','What to check'],[
('Region','The source must identify NCR records. Verify the region code using that survey’s dictionary.'),
('Annual family income','Total annual income for each family, in pesos; not per-person income.'),
('Annual family expenditure','Annual expenditure for the same family, in pesos.'),
('Family size','The number of family members represented by that record.'),
('Survey weight','The valid positive family survey expansion weight for that survey year.'),
('Official values entered separately','NCR poverty threshold per person per year and NCR poverty incidence among families for the matching year.')],[1.7,4.9])
h('Check column names')
p('The processor recognizes survey-specific aliases. Examples include TOTINC for income; TOTEXP, TTOTEX and T_TOTEX for expenditure; FSIZE for family size; and RFACT or RFADJ for survey weight. These are examples, not interchangeable definitions. Consult the source dictionary before renaming any column.')
h('Check the threshold units')
p('Enter the annual threshold for one person. The system multiplies that value by each family’s size before comparing it with annual family income. Do not enter a five-person family threshold as a per-person threshold, and do not multiply by family size yourself before entry.')
p('Illustrative example: an annual per-person threshold of PHP 30,000 and a family size of 4 gives a family income threshold of PHP 120,000. An annual family income of PHP 110,000 is below that threshold.')
note('If your official source gives a monthly amount, establish the equivalent annual per-person amount before entry. Record the conversion and source. Poverty incidence is entered as a percentage such as 2.50, not 0.025.')

page(4,'Upload one or many datasets')
h('Upload one file')
steps(['Open Dataset Management.','Click Add Dataset. Select the CSV, XLSX or XLS file and confirm the file selection.','Remain on the page while the upload runs. A full progress bar may still be followed by saving on the server.','Confirm the file appears under Pending Uploads. Check the filename before processing it.'])
h('Bulk upload')
steps(['Click Bulk Upload. Select several supported files in the file picker.','Review the upload results as the queue runs. The files are uploaded in sequence.','When the queue finishes, identify which files succeeded, failed or remain unconfirmed.','Check Pending Uploads before retrying an unconfirmed file so that you do not create an unnecessary duplicate.','Process each uploaded file separately with its own survey year, threshold and incidence. Bulk upload does not apply one set of processing values to all files.'])
h('Understand the queue results')
table(['Result','What to do'],[
('Uploaded','The upload was confirmed. Continue to processing.'),
('Failed','Read the reported error, correct it and retry that file.'),
('Unconfirmed','The response was uncertain. Check whether the file was saved before retrying.'),
('Stopped or not uploaded','The queue ended before this file was confirmed. Review the saved files before restarting.')],[1.6,5.0])
h('Find an uploaded file')
p('Use the filename search and the From and To dates. These dates filter upload dates in Dataset Management. Use pagination to see additional records; the pending and processed lists have separate pages.')
note('If a selected file has an unsupported extension, the upload control rejects the selection. Changing an extension alone does not convert a file into a valid CSV or Excel workbook.')

page(5,'Process a dataset')
h('Enter the processing information')
steps(['Find the file under Pending Uploads and open its Process action.','Confirm the filename in the Process Dataset window. Survey Type is FIES.','Enter the survey Year. This is the year represented by the data, not the upload date.','Enter Poverty Threshold (PHP) as the positive annual NCR threshold per person.','Enter Poverty Incidence (%) for NCR families in the same year, from the official source.','Click Process. Stay on the page and wait for the completion message.','Review the record under Processed Datasets before merging it into the history.'])
h('What processing does')
p('The processor identifies the required columns and NCR records, checks usable numeric values, and selects families whose annual income is below their family threshold. The family threshold is the entered annual per-person threshold multiplied by that family’s size.')
p('It calculates survey-weighted mean income, expenditure and family size for the selected families. The same selected group supplies all three indicators. It stores the poverty incidence that you entered as the target for that year.')
p('Weighted mean = sum of each selected family’s value multiplied by its survey weight, divided by the sum of those survey weights.')
h('Check the output')
bullets(['Confirm the region and survey year.','Check that the displayed income and expenditure are plausible annual averages for the selected families.','Keep average family size as a decimal. It is a mean, so a value such as 4.27 is valid.','Compare the saved poverty incidence with the value you entered.'])
note('The displayed values use two decimal places by truncation. For example, 4.279 displays as 4.27. Display formatting does not replace the underlying numeric precision. A missing-column error must be resolved using the survey dictionary; do not fill it with an invented value.')

page(6,'Merge processed datasets')
p('Merging places processed values into the historical data used by the pipeline. Uploading or processing alone does not complete this step. Review the year, threshold basis and poverty incidence before merging.')
h('Merge one dataset')
steps(['In Processed Datasets, choose the Merge action for the approved file.','Review the region, year and poverty incidence in Merge into Actuals?','Click Merge to proceed, or Cancel to keep the current history unchanged.','Wait for the success message. Open Data and verify the resulting year and values.','Rerun forecasting and model training when the history changes.'])
h('Bulk merge')
steps(['Select the checkboxes beside the processed datasets you intend to merge.','Use the table header checkbox for the current page, or Select all filtered to include all matching records. Use Clear selection to start again.','Choose only one dataset for each region and year. Resolve duplicate-year selections before proceeding.','Click Merge Selected. Review the selected datasets and the update warning in the confirmation window.','Confirm Merge Selected. Wait for the queue and inspect the result for every file.','Verify the merged years in Data, then rerun forecasting and training.'])
h('Understand updates and partial completion')
p('When the same region and year already exist, merging updates those actual values. Check the selected file carefully: a successful merge can replace a previous value for that year.')
p('Bulk merges run separately for each dataset. If you stop the operation or one item fails, merges already completed remain saved. Review the per-file results; do not assume the entire batch was rolled back.')
h('Correct an incorrect entry')
p('Prepare and process a corrected copy with the right year, threshold and poverty incidence. Review its output and merge it into the matching year. Keep a record of the correction. Use deletion only when removing a record is actually intended.')

page(7,'Review data and interpolation')
h('Inspect the history')
steps(['Open Data.','Search by region or year. The From and To controls on this page filter the historical year, unlike the upload-date filters in Dataset Management.','Use pagination to inspect all matching records.','Check the source of each row and distinguish observed FIES values from interpolated values.'])
h('Understand interpolation')
p('Interpolation estimates missing years between usable observations. For example, when actual poverty incidence exists for 2023 and 2025, the missing 2024 value can be filled from those surrounding years. It is an estimate, not a newly observed survey result.')
p('Illustrative example: if incidence is 3.00% in 2023 and 2.00% in 2025, linear interpolation gives 2.50% for 2024. This example assumes compatible observations and definitions.')
p('In the active FIES-only history, interpolated rows are calculated from the surrounding FIES observations. Correct the underlying observations when an interpolated value is wrong. The page does not allow you to independently delete those calculated rows.')
h('When a new survey year arrives')
steps(['Upload and process the new FIES file using its own official values.','Merge the new record and review the updated history.','Check that eligible gaps between observations now show interpolated values.','Run Forecast and Train Models again so results reflect the changed history.'])
h('Delete only an intended actual row')
p('For a source row that must be removed, use its delete action in Data. Check the region and year in Delete Data Row? and choose Delete only if the removal is intended. Cancel closes the confirmation without deleting it. Rerun the pipeline afterward.')
note('Deleting an upload in Dataset Management removes that uploaded file and dataset record. Removing a historical row in Data is a different operation. Verify Data separately; do not assume deleting an upload reverses an earlier merge.')

page(8,'Run the three year forecast')
p('Forecasting estimates the three socioeconomic inputs for the next three years. Poverty incidence is predicted later by the selected machine learning model using those forecast inputs.')
h('Generate and inspect the forecast')
steps(['Finish merging the intended historical data and review it in Data.','Open Forecasting and click Run Forecast.','Remain on the page while the operation runs.','Check Forecasted Values for the three expected years and the income, expenditure and family-size estimates.','In the trend chart, select an indicator and hover over points. Read the legend to distinguish actual, interpolated and forecast values.','Proceed to Model Training to generate poverty predictions from the current data.'])
h('How the forecast window moves')
table(['Latest eligible history year','Next forecast and prediction years'],[
('2023','2024, 2025 and 2026'),
('2026 after a new actual release','2027, 2028 and 2029'),
('2022 after removing the latest 2023 observation','2023, 2024 and 2025')],[3.2,3.4])
p('Running the forecast repeatedly without adding a newer actual or interpolated year does not keep extending the horizon. The active result window stays three years after the latest eligible historical year.')
h('Retained forecasts and new observations')
p('When a later actual record arrives, earlier forecast inputs can remain as historical bridge values where no eligible actual or interpolated inputs exist. For example, previous 2024 and 2025 forecast inputs can support the history when a 2026 observation is added. Actual or eligible interpolated values take precedence when available.')
note('ARIMA forecasts may be flat. This alone does not mean the run failed. Inspect the source history and units before changing data. Forecast values and retained forecast inputs must not be described as official observations.')

page(9,'Train models and predict poverty')
h('Run training')
steps(['Confirm the intended datasets have been merged and the historical values are correct.','Open Model, then Model Training.','Click Train Models once. Wait while the page reports loading, preprocessing, training, evaluation and prediction progress.','Stay on the page until Training complete and the best model name appear.','Review the Predicted Poverty Incidence by Year chart and Predictions table. Confirm the region, three-year window and model used.','Open Model Evaluation, then Analysis, to review how the result was obtained.'])
h('What happens inside the pipeline')
table(['Stage','Current behavior'],[
('Prepare history','Use actual and interpolated FIES history for the poverty model.'),
('Prepare features','Use income, expenditure and family size as the three configured inputs.'),
('Validate and tune','Use shuffled five-fold cross-validation, with tuning inside training folds.'),
('Compare models','Evaluate XGBoost, Ridge and Random Forest. Select using RMSE.'),
('Generate forecasts','The complete pipeline refreshes socioeconomic forecasts for the three-year window.'),
('Predict poverty','Apply the selected model to forecast inputs and save the resulting poverty estimates.')],[1.55,5.05])
p('Predicted poverty incidence is not used as a training target. Actual and interpolated poverty incidence form the target history. Retained forecast indicators used by the forecasting history are distinct from the poverty model’s target data.')
h('Check that the run is current')
p('A previous successful run may still be displayed when a new request fails. Read any new error message and verify that the latest run completed. After changing historical data, do not treat old saved predictions as refreshed results until training completes.')
note('Do not add artificial changes to the inputs just to make predictions differ by year. Equal values can arise when the forecast inputs are similar, when tree models return the same leaf output, or when display truncation hides small differences.')

page(10,'Read the model evaluation')
steps(['Open Model, then Model Evaluation after a completed training run.','Compare the Metrics Comparison rows for XGBoost, Ridge and Random Forest.','Identify the Best Model label and review RMSE, MAE and R² together.','Use Visual Comparison to compare error metrics and R² across models. Record the run context when discussing the result.'])
table(['Metric','How to interpret it'],[
('RMSE','Lower is better. It gives more weight to large errors. The current configuration uses RMSE to select the best model.'),
('MAE','Lower is better. It is the average absolute prediction error. For a percentage target, interpret the error in percentage points.'),
('R²','Higher is generally better. A negative value means the model performed worse than the evaluation-mean reference underlying R² for those targets.'),
('Best Model','Best among the compared models under the configured evaluation. It does not guarantee accurate future predictions.')],[1.3,5.3])
h('Understand the current validation')
p('The active configuration uses shuffled five-fold cross-validation and includes actual and interpolated target rows. Each outer fold is evaluated using rows withheld from that fold’s training; tuning happens within the training portion.')
p('This evaluates the model on the available indicator rows. It does not directly measure the accuracy of forecasting three future years. Shuffling years and including interpolated targets can produce more optimistic scores than evaluating only future observed years in chronological order.')
h('Discuss results accurately')
bullets(['State that the current scores come from shuffled K-fold validation.','Report the data period and distinguish actual observations from interpolated years.','Compare RMSE and MAE alongside R²; a small sample can make R² unstable.','Treat the three-year predictions as estimates for analysis, not official releases.'])
note('A better K-fold score does not prove that a change improved real future forecasting. Compare equivalent validation setups and ask the maintainer for a chronological backtest when future forecasting accuracy is the question.')

page(11,'Analysis and visualization')
h('Review Analysis')
steps(['Open Analysis after a successful model run.','Read Feature Importance to see the relative influence of the available inputs in the saved model explanation.','Review Prediction Trend to compare historical poverty incidence with the saved future estimates.','Hover over chart points to inspect their values and years.'])
p('Feature importance describes how the model uses the inputs. It does not prove that changing a variable will cause poverty to rise or fall. The direction and meaning of an association require additional analysis.')
h('Review Visualization')
steps(['Open Visualization. Review the summary cards and poverty timeline.','Choose Show Indicator Trends to inspect changes in the input indicators over time.','Choose Year-over-Year Change to inspect changes in poverty incidence relative to the prior year.','Choose Indicator Correlation to explore the relationship between an indicator and poverty incidence.','Use available chart selectors, legends and tooltips to identify the series and year being compared. Close the chart window when finished.'])
h('Read charts consistently')
bullets(['Distinguish actual, interpolated and forecast or predicted series using the legend.','A connected line helps show continuity; it does not turn an interpolated point into an observed value.','Compare years on the same basis. A change in the surveyed family group or threshold units can distort comparisons.','Income and expenditure are in pesos, family size is a mean number of members, and poverty incidence is a percentage.'])
h('If a caption looks outdated')
p('Some visualization captions retain older wording about six indicators or a fixed historical date range. The active pipeline uses three inputs. Use the actual plotted series, table years and current configuration when interpreting results, and report misleading captions to the maintainer.')
note('When presenting the system, show Data first, then Forecasting and Model Training results. This makes it clear which values were observed, interpolated, forecast or predicted.')

page(12,'Manage users and audit logs')
p('The System menu is shown to super administrators. Ordinary administrators use the data and model modules; request help from a super administrator when account administration is needed.')
h('Create an account')
steps(['Open System, then User Management. Click Add User.','Enter Name, Username, Email, Address and Password. Choose Admin or Super Admin according to the access the person needs.','Follow the displayed password rule: at least eight characters, with uppercase, lowercase and a special character.','Click Create User and confirm the account appears in the list. Share credentials privately.'])
h('Edit or reset another account')
p('Use search by name, username or email and the Role filter to locate the person. Choose Edit, update Name, Email or Address, and click Save Changes. The edit form does not change the username or role.')
p('To reset another user’s password, open Change Password, enter and confirm the new password, and submit the form. A self-password change requires the current password on the backend; the administrative reset form does not collect it. If your own reset is rejected, ask another authorized super administrator or the maintainer for assistance.')
h('Remove an account only when intended')
p('Choose the user’s delete action, review the identity in Delete User?, and confirm Delete only when authorized. Cancel leaves the account in place. Do not delete working accounts as part of an evaluation demonstration.')
h('Review Audit Logs')
steps(['Open System, then Audit Logs.','Review the user, action, details and time recorded for an operation.','Use pagination to move through the history. The page loads ten records at a time.','Use search or date filters to narrow the currently loaded page. Move to another page to check additional records.'])
note('The total count represents the stored history. Search and date filters currently act on the loaded page, so an empty filtered page does not prove that no matching event exists elsewhere. Pagination does not delete older logs.')

page(13,'Cancel work and resolve errors')
h('Leave during an operation')
p('Upload, preprocessing, merging, forecasting and training use a leave-page warning. When Leave this page? appears, choose No to stay and let the operation continue. Choose Yes to request cancellation and leave after the application handles the stop request.')
p('Wait while Stopping is displayed. If cancellation fails, remain on the page and read the error. Refreshing or closing a browser uses the browser’s own warning; it may not show the same wording. Already completed uploads or merges remain saved. Cancellation is not an undo action.')
table(['Problem','Recommended response'],[
('Cannot find a required column','Check the survey dictionary and header row. Confirm the column represents the requested quantity. Ask the maintainer to add a verified alias if necessary.'),
('No usable families or invalid values','Check region codes, numeric fields, positive weights, family size and annual per-person threshold units.'),
('Failed to fetch or Network Error','Ask the operator to verify Docker and ngrok. Do not repeatedly submit the same upload or merge before checking saved results.'),
('ERR_NGROK_3200','The public tunnel is offline. The operator must restart ngrok and confirm its forwarding address.'),
('No predictions or evaluation results','Verify historical data has been merged, then complete Train Models.'),
('Identical yearly predictions','Compare the forecast inputs, latest successful run and display precision. Similar inputs or model behavior can legitimately give equal outputs.'),
('Unexpected forecast years','Check the latest actual or interpolated year in Data, then rerun Forecast and Train Models.'),
('Training failed','Read the reported error, confirm sufficient usable historical rows and consistent units, and contact the maintainer if it persists.')],[2.0,4.6])
note('When reporting a failure, provide the module, action, time, filename and survey year where relevant, and the exact error. Exclude passwords, authentication tokens and raw household records from screenshots.')

page(14,'Start and stop evaluator access')
p('This section is for the operator on the computer containing the project and Docker data. Evaluators only need a browser, the public frontend link and an assigned account.')
h('Start the local services')
steps(['Open Docker Desktop and wait for it to start.','If the existing project is stopped, start its group in Docker Desktop. Expand it and confirm the database, ML, backend and frontend services are running.','Alternatively, open PowerShell and run the following commands, one line at a time.'])
code('cd C:\\Users\\Dell\\Desktop\\Capstone\\poverty-prediction-system\ndocker compose --env-file .env.docker up -d')
p('Open http://localhost:8080 and verify that the system loads. If the local system fails, resolve that before starting remote access.')
h('Start the tunnel')
p('In a separate terminal with ngrok configured, run:')
code('ngrok http 8080')
p('Keep the terminal open. Confirm Session Status is online and the forwarding address points to localhost:8080. The current frontend configuration expects https://county-ladder-kisser.ngrok-free.dev. If the address changes, follow section 15 before sharing the site.')
h('Check the evaluator experience')
steps(['Open https://poverty-prediction-system-five.vercel.app in a browser.','Sign in with an assigned account and confirm Data or the Dashboard loads its records.','Check one results page. A visible login screen alone does not verify the backend connection.','Keep the computer awake, connected to the internet, and running Docker and ngrok throughout the evaluation.'])
h('Stop after the session')
p('Wait for active work to finish. Press Ctrl+C in the ngrok terminal to stop public access. To stop the application services while preserving their stored data, use Docker Desktop’s stop action or run:')
code('docker compose --env-file .env.docker stop')

page(15,'Maintain deployment and preserve data')
h('Understand the connection')
p('The evaluator opens the frontend on Vercel. Requests under /api are forwarded through ngrok to the local Docker application. The backend, machine learning service, database, uploaded datasets and saved models remain on the operator computer.')
h('When code changes')
table(['Changed part','Operator action'],[
('Frontend','Deploy the updated frontend to Vercel. Rebuild the Docker frontend as well if the local interface needs the same change.'),
('Backend or ML','Rebuild and restart the corresponding Docker service. A Vercel deployment does not update local backend or ML code.'),
('ngrok address','Update the /api destination in frontend/vercel.json to the active HTTPS tunnel address, then redeploy the frontend.'),
('Source data','Use the application’s upload, process and merge workflow. Rerun forecasting and training; a frontend redeployment is not required.')],[1.5,5.1])
p('The Vercel project uses the frontend directory, Vite, npm run build and the dist output directory. Keep API requests relative to the site through /api. Do not set a public frontend API address to localhost: that would refer to each evaluator’s own device.')
h('Keep recoverable copies')
bullets(['Keep original FIES files, their dictionaries and the official threshold and incidence references outside the application.','Before maintenance, ask the operator or maintainer to take a current database backup and copies of the uploaded-data and saved-model folders.','The Docker database volume, docker-state/ml-data and docker-state/saved-models contain persistent application state. Preserve them.','An initial database import file is not automatically a current backup. Verify backups and restoration on a separate copy before relying on them.','Do not remove Docker volumes, use factory reset, or delete data folders as a routine restart procedure.'])
h('End of session checklist')
bullets(['Confirm any uploads, merges and training runs completed or were deliberately cancelled.','Check that displayed predictions correspond to the intended history and three-year window.','Log out of shared browsers.','Stop the tunnel when evaluator access is no longer needed. Keep application data and backups intact.'])

out=ROOT/'Poverty Prediction System User Manual.docx'
doc.save(out)
print(out)
